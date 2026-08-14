from pathlib import Path

from openpyxl import load_workbook

from core.excel_inspector import ExcelInspector
from core.remediation_applier import apply_remediation, resolve_affected_cells
from core.iteration_lineage import latest_iteration, next_iteration

FIXTURE = Path(__file__).parent / "fixtures" / "test_cases" / "t101_fail_all.xlsx"


def test_legacy_finding_locations_are_recovered_for_display():
    finding = next(f for f in ExcelInspector().inspect_workbook(str(FIXTURE)) if f.rule_id == "T101-NO-EMPTY-CELLS")
    finding.affected_cells = []
    assert resolve_affected_cells(str(FIXTURE), finding) == ["B2", "B3", "B5"]


def test_symbol_iteration_has_no_empty_data_cells(tmp_path):
    finding = next(f for f in ExcelInspector().inspect_workbook(str(FIXTURE)) if f.rule_id == "T101-NO-EMPTY-CELLS")
    target = tmp_path / "iteration.xlsx"
    result = apply_remediation(str(FIXTURE), str(target), finding, "select_symbol", "r", "B2")

    assert result.success
    remaining = [f for f in ExcelInspector().inspect_workbook(str(target)) if f.rule_id == "T101-NO-EMPTY-CELLS"]
    assert any("B3" in f.location and "B5" in f.location for f in remaining)
    ws = load_workbook(target)["Data"]
    assert ws["B2"].value == "r"
    assert ws["B3"].value is None
    assert ws["B5"].value is None


def test_gridline_iteration_is_revalidated(tmp_path):
    finding = next(f for f in ExcelInspector().inspect_workbook(str(FIXTURE)) if f.rule_id == "T101-GRIDLINES-ON")
    target = tmp_path / "iteration.xlsx"
    result = apply_remediation(str(FIXTURE), str(target), finding, "turn_on", None)

    assert result.success
    assert not any(f.rule_id == "T101-GRIDLINES-ON" and f.sheet_name == "Data" for f in ExcelInspector().inspect_workbook(str(target)))
    assert load_workbook(target)["Data"].sheet_view.showGridLines is True


def test_subsequent_cell_iteration_is_cumulative_and_not_overwritten(tmp_path):
    finding = next(f for f in ExcelInspector().inspect_workbook(str(FIXTURE)) if f.rule_id == "T101-NO-EMPTY-CELLS")
    first = tmp_path / "iteration_01_FND-ONE_t101_fail_all.xlsx"
    second = tmp_path / "iteration_02_FND-TWO_t101_fail_all.xlsx"

    assert apply_remediation(str(FIXTURE), str(first), finding, "select_symbol", "r", "B2").success
    current = next(f for f in ExcelInspector().inspect_workbook(str(first)) if f.rule_id == "T101-NO-EMPTY-CELLS")
    assert apply_remediation(str(first), str(second), current, "select_symbol", "r", "B3").success

    ws = load_workbook(second)["Data"]
    assert ws["B2"].value == "r"
    assert ws["B3"].value == "r"
    assert ws["B5"].value is None
    assert first.exists()
    assert load_workbook(first)["Data"]["B3"].value is None


def test_second_iteration_rejects_cell_fixed_in_first_iteration(tmp_path):
    finding = next(f for f in ExcelInspector().inspect_workbook(str(FIXTURE)) if f.rule_id == "T101-NO-EMPTY-CELLS")
    first = tmp_path / "iteration_01_FND-ONE_t101_fail_all.xlsx"
    second = tmp_path / "iteration_02_FND-TWO_t101_fail_all.xlsx"
    assert apply_remediation(str(FIXTURE), str(first), finding, "select_symbol", "r", "B2").success

    # Deliberately use the stale finding from the initial revision.
    result = apply_remediation(str(first), str(second), finding, "select_symbol", "x", "B2")
    assert result.success is False
    assert "not an affected cell" in result.message
    assert not second.exists()


def test_iteration_lineage_chooses_latest_and_allocates_unique_name(tmp_path):
    original = "t101_fail_all.xlsx"
    first = tmp_path / "iteration_01_FND-ONE_t101_fail_all.xlsx"
    first.write_bytes(b"one")
    assert latest_iteration(tmp_path, original) == first
    next_path = next_iteration(tmp_path, original, "FND-TWO")
    assert next_path.name == "iteration_02_FND-TWO_t101_fail_all.xlsx"
