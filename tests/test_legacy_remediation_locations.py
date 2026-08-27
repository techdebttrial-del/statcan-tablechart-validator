from pathlib import Path

from core.excel_inspector import ExcelInspector
from core.remediation_applier import apply_remediation
from core.models import Finding

FIXTURE = Path(__file__).parent / "fixtures" / "test_cases" / "t101_fail_all.xlsx"


def test_legacy_finding_can_be_resolved_against_current_workbook(tmp_path):
    findings = ExcelInspector().inspect_workbook(str(FIXTURE))
    current = next(f for f in findings if f.rule_id == "T101-NO-EMPTY-CELLS")
    legacy = Finding(
        finding_id=current.finding_id, rule_id=current.rule_id, pack=current.pack,
        severity=current.severity, sheet_name=current.sheet_name,
        location="2 empty cell(s)", title_en=current.title_en,
        title_fr=current.title_fr, description_en=current.description_en,
        description_fr=current.description_fr,
    )
    target = tmp_path / "iteration.xlsx"

    result = apply_remediation(str(FIXTURE), str(target), legacy, "select_symbol", "..")

    assert result.success is True
    from openpyxl import load_workbook
    ws = load_workbook(target)["Data"]
    assert [ws[c].value for c in current.affected_cells] == ["..", ".."]


def test_legacy_formula_finding_is_resolved_before_write(tmp_path):
    findings = ExcelInspector().inspect_workbook(str(FIXTURE))
    current = next(f for f in findings if f.rule_id == "T101-NO-FORMULAS")
    current.location = "1 formula cell(s)"
    current.affected_cells = []
    target = tmp_path / "iteration.xlsx"

    result = apply_remediation(str(FIXTURE), str(target), current, "replace_with_value", "42")

    assert result.success is True
    from openpyxl import load_workbook
    # Formula is at A2 in the new fixture (A2 is part of merged A2:A3)
    assert load_workbook(target, data_only=False)["Data"]["A2"].value == 42
