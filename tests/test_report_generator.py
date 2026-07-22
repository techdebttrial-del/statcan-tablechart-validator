"""
Tests for ReportGenerator — bilingual output, finding tables, decision rendering.
"""
import os
import sys
import tempfile

import pytest
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.repository_client import InMemoryRepositoryClient
from core.project_store import ProjectStore
from core.report_generator import ReportGenerator
from core.translation_manager import TranslationManager, PassthroughTranslator
from core.models import DecisionType, DecisionReason, ComplianceStatus


def _make_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "tbl1"
    ws["A1"] = "Category"
    ws["B1"] = "Value"
    ws["A2"] = "Row1"
    ws["B2"] = 100
    ws["A3"] = "Source: Test"
    ws["A4"] = "x = suppressed"
    ws["A5"] = "Unit of measure: thousands"
    wb.save(path)


def _make_project_with_revision():
    store = ProjectStore(repo=InMemoryRepositoryClient())
    p = store.create_project(label="Test Report Project", reviewer="Dr. Reviewer")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        _make_workbook(f.name)
        rev = store.add_workbook_revision(p, f.name, "test.xlsx", "Dr. Reviewer")
        os.unlink(f.name)
    return store, p, rev


class TestReportGenerator:
    def test_generate_both_returns_en_and_fr(self):
        _, p, _ = _make_project_with_revision()
        gen = ReportGenerator()
        reports = gen.generate_both(p)
        assert "en" in reports
        assert "fr" in reports
        assert len(reports["en"]) > 0
        assert len(reports["fr"]) > 0

    def test_en_report_contains_project_info(self):
        _, p, _ = _make_project_with_revision()
        gen = ReportGenerator()
        report = gen.generate(p, "en")
        assert "Test Report Project" in report
        assert "Dr. Reviewer" in report
        assert p.project_id in report

    def test_fr_report_contains_project_info(self):
        _, p, _ = _make_project_with_revision()
        gen = ReportGenerator()
        report = gen.generate(p, "fr")
        assert "Test Report Project" in report
        assert "Dr. Reviewer" in report

    def test_en_report_uses_english_strings(self):
        _, p, _ = _make_project_with_revision()
        gen = ReportGenerator()
        report = gen.generate(p, "en")
        assert "Table 101 / Chart 101 Compliance Review" in report
        assert "Findings" in report

    def test_fr_report_uses_french_strings(self):
        _, p, _ = _make_project_with_revision()
        gen = ReportGenerator()
        report = gen.generate(p, "fr")
        assert "Examen de conformité" in report
        assert "Constats" in report

    def test_report_includes_findings_table(self):
        _, p, _ = _make_project_with_revision()
        gen = ReportGenerator()
        report = gen.generate(p, "en")
        # Should have a table with rule IDs
        assert "T101-" in report or "No findings" in report

    def test_report_shows_compliance_status(self):
        _, p, rev = _make_project_with_revision()
        gen = ReportGenerator()
        report = gen.generate(p, "en")
        if rev.compliance_status == ComplianceStatus.NOT_COMPLIANT:
            assert "Not compliant" in report
        else:
            assert "Compliant" in report

    def test_report_includes_decisions_when_present(self):
        store, p, rev = _make_project_with_revision()
        if rev.findings:
            store.record_decision(
                p, rev, rev.findings[0].finding_id,
                decision_type=DecisionType.WAIVED,
                reason=DecisionReason.EXCEPTION_APPROVED,
                note_text="Approved for this publication cycle.",
                note_language="en",
                reviewer="Dr. Reviewer",
            )
        gen = ReportGenerator()
        report = gen.generate(p, "en")
        assert "waived" in report.lower() or "No decisions" in report

    def test_report_includes_audit_trail(self):
        _, p, _ = _make_project_with_revision()
        gen = ReportGenerator()
        report = gen.generate(p, "en")
        assert "Audit Trail" in report or "Amendments" in report

    def test_empty_project_report(self):
        store = ProjectStore(repo=InMemoryRepositoryClient())
        p = store.create_project(label="Empty", reviewer="A")
        gen = ReportGenerator()
        report = gen.generate(p, "en")
        assert "Empty" in report
        assert "A" in report
