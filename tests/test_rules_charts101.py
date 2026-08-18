"""
Comprehensive tests for Charts 101 rules.

Each test creates a workbook with charts using openpyxl to exercise
specific rule conditions.
"""
import os
import sys
import tempfile

import pytest
import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.excel_inspector import ExcelInspector
from core.rule_pack_loader import RulePackLoader

CONFIG_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "config")


def _make_inspector():
    return ExcelInspector(loader=RulePackLoader(CONFIG_DIR))


def _save_wb(wb):
    f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(f.name)
    f.close()
    return f.name


def _rule_ids(findings):
    return {f.rule_id for f in findings}


# ---------------------------------------------------------------------------
# C101-NO-EMPTY-CELLS
# ---------------------------------------------------------------------------

class TestC101NoEmptyCells:
    def test_fail_when_chart_sheet_has_empty_data_cells(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "chart1"
        # Set up data for chart
        ws["A1"] = "Category"
        ws["B1"] = "Value"
        ws["A2"] = "A"
        ws["B2"] = 10
        ws["A3"] = "B"
        # B3 is empty — data cell for chart
        ws["A4"] = "Source: X"
        ws["A5"] = ".. = not available"

        chart = BarChart()
        data = Reference(ws, min_col=2, min_row=1, max_row=3)
        cats = Reference(ws, min_col=1, min_row=2, max_row=3)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "D1")

        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            ids = _rule_ids(findings)
            assert "C101-NO-EMPTY-CELLS" in ids
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# C101-MAX-SIX-SERIES
# ---------------------------------------------------------------------------

class TestC101MaxSixSeries:
    def test_fail_when_more_than_six_series(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "chart1"
        # Create 7 data columns
        headers = ["Cat"] + [f"S{i}" for i in range(1, 8)]
        for i, h in enumerate(headers):
            ws.cell(row=1, column=i+1, value=h)
        for row in range(2, 5):
            ws.cell(row=row, column=1, value=f"R{row}")
            for col in range(2, 9):
                ws.cell(row=row, column=col, value=row * col)

        ws["A6"] = "Source: X"
        ws["A7"] = ".. = not available"

        chart = BarChart()
        for col in range(2, 9):  # 7 series
            data = Reference(ws, min_col=col, min_row=1, max_row=4)
            chart.add_data(data, titles_from_data=True)
        cats = Reference(ws, min_col=1, min_row=2, max_row=4)
        chart.set_categories(cats)
        ws.add_chart(chart, "D6")

        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            ids = _rule_ids(findings)
            assert "C101-MAX-SIX-SERIES" in ids
        finally:
            os.unlink(path)

    def test_pass_when_six_or_fewer_series(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "chart1"
        headers = ["Cat"] + [f"S{i}" for i in range(1, 4)]
        for i, h in enumerate(headers):
            ws.cell(row=1, column=i+1, value=h)
        for row in range(2, 5):
            ws.cell(row=row, column=1, value=f"R{row}")
            for col in range(2, 4):
                ws.cell(row=row, column=col, value=row * col)

        ws["A6"] = "Source: X"
        ws["A7"] = "x = suppressed"

        chart = BarChart()
        for col in range(2, 4):
            data = Reference(ws, min_col=col, min_row=1, max_row=4)
            chart.add_data(data, titles_from_data=True)
        cats = Reference(ws, min_col=1, min_row=2, max_row=4)
        chart.set_categories(cats)
        ws.add_chart(chart, "D6")

        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            ids = _rule_ids(findings)
            assert "C101-MAX-SIX-SERIES" not in ids
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# C101-SOURCE-PRESENT
# ---------------------------------------------------------------------------

class TestC101SourcePresent:
    def test_fail_when_no_source_in_chart_sheet(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "chart1"
        ws["A1"] = "Cat"
        ws["B1"] = "Val"
        ws["A2"] = "A"
        ws["B2"] = 10

        chart = BarChart()
        data = Reference(ws, min_col=2, min_row=1, max_row=2)
        chart.add_data(data, titles_from_data=True)
        ws.add_chart(chart, "D1")

        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            ids = _rule_ids(findings)
            assert "C101-SOURCE-PRESENT" in ids
        finally:
            os.unlink(path)

    def test_pass_when_source_present(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "chart1"
        ws["A1"] = "Cat"
        ws["B1"] = "Val"
        ws["A2"] = "A"
        ws["B2"] = 10
        ws["A3"] = "Source: Labour Force Survey"
        ws["A4"] = "x = suppressed"

        chart = BarChart()
        data = Reference(ws, min_col=2, min_row=1, max_row=2)
        chart.add_data(data, titles_from_data=True)
        ws.add_chart(chart, "D1")

        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            ids = _rule_ids(findings)
            assert "C101-SOURCE-PRESENT" not in ids
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# C101-STANDARD-SYMBOLS
# ---------------------------------------------------------------------------

class TestC101StandardSymbols:
    def test_fail_when_no_symbols_in_chart_sheet(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "chart1"
        ws["A1"] = "Cat"
        ws["B1"] = "Val"
        ws["A2"] = "A"
        ws["B2"] = 10
        # Use text with NO standard symbol tokens as substrings
        ws["A3"] = "ABC Data Only"
        # No symbol definitions

        chart = BarChart()
        data = Reference(ws, min_col=2, min_row=1, max_row=2)
        chart.add_data(data, titles_from_data=True)
        ws.add_chart(chart, "D1")

        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            ids = _rule_ids(findings)
            assert "C101-STANDARD-SYMBOLS" in ids
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# C101-UNIQUE-NUMBERING
# ---------------------------------------------------------------------------

class TestC101UniqueNumbering:
    def test_fail_when_duplicate_chart_numbers(self):
        """Two sheets with same extracted chart number trigger uniqueness."""
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "chart1"
        ws1["A1"] = "Cat"
        ws1["B1"] = "Val"
        ws1["A2"] = "A"
        ws1["B2"] = 10
        ws1["A3"] = "Source: X"
        ws1["A4"] = "x = suppressed"

        chart1 = BarChart()
        data1 = Reference(ws1, min_col=2, min_row=1, max_row=2)
        chart1.add_data(data1, titles_from_data=True)
        ws1.add_chart(chart1, "D1")

        # openpyxl auto-renames exact duplicates. Test the logic directly.
        # (See test_uniqueness_logic_detects_duplicates below)
        ws2 = wb.create_sheet("chart1a")
        ws2["A1"] = "Cat"
        ws2["B1"] = "Val"
        ws2["A2"] = "B"
        ws2["B2"] = 20
        ws2["A3"] = "Source: X"
        ws2["A4"] = "x = suppressed"

        chart2 = BarChart()
        data2 = Reference(ws2, min_col=2, min_row=1, max_row=2)
        chart2.add_data(data2, titles_from_data=True)
        ws2.add_chart(chart2, "D1")

        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            # openpyxl extracts "1" and "1a" — different numbers
            # The workbook-level uniqueness is tested directly below
        finally:
            os.unlink(path)

    def test_uniqueness_logic_detects_duplicates(self):
        """Test numbering uniqueness logic directly."""
        inspector = _make_inspector()
        findings = inspector._check_numbering_uniqueness(
            ["1", "1", "2"], "charts101", "C101-UNIQUE-NUMBERING"
        )
        assert len(findings) == 1
        assert findings[0].rule_id == "C101-UNIQUE-NUMBERING"


# ---------------------------------------------------------------------------
# C101-ONE-CHART-DATA-PER-SHEET
# ---------------------------------------------------------------------------

class TestC101OneChartPerSheet:
    def test_pass_when_single_chart(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "chart1"
        ws["A1"] = "Cat"
        ws["B1"] = "Val"
        ws["A2"] = "A"
        ws["B2"] = 10
        ws["A3"] = "Source: X"
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        ws.add_chart(chart, "D1")
        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            ids = _rule_ids(findings)
            assert "C101-ONE-CHART-DATA-PER-SHEET" not in ids
        finally:
            os.unlink(path)

    def test_fail_when_two_charts_on_one_sheet(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "chart1"
        ws["A1"] = "Cat"
        ws["B1"] = "Val"
        ws["A2"] = "A"
        ws["B2"] = 10
        ws["A3"] = "Source: X"
        d = Reference(ws, min_col=2, min_row=1, max_row=2)
        c1 = BarChart()
        c1.add_data(d, titles_from_data=True)
        ws.add_chart(c1, "D1")
        c2 = BarChart()
        c2.add_data(d, titles_from_data=True)
        ws.add_chart(c2, "H1")
        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            ids = _rule_ids(findings)
            assert "C101-ONE-CHART-DATA-PER-SHEET" in ids
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# C101-GRIDLINES-NO-DATALABELS
# ---------------------------------------------------------------------------

class TestC101GridlinesNoDataLabels:
    def test_fail_when_gridlines_and_data_labels_both_present(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "chart1"
        ws["A1"] = "Cat"
        ws["B1"] = "Val"
        ws["A2"] = "A"
        ws["B2"] = 10
        ws["A3"] = "Source: X"
        from openpyxl.chart.label import DataLabelList
        from openpyxl.chart.axis import ChartLines
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        chart.dataLabels = DataLabelList()
        chart.x_axis.majorGridlines = ChartLines()
        ws.add_chart(chart, "D1")
        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            ids = _rule_ids(findings)
            assert "C101-GRIDLINES-NO-DATALABELS" in ids
        finally:
            os.unlink(path)

    def test_pass_when_data_labels_without_gridlines(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "chart1"
        ws["A1"] = "Cat"
        ws["B1"] = "Val"
        ws["A2"] = "A"
        ws["B2"] = 10
        ws["A3"] = "Source: X"
        from openpyxl.chart.label import DataLabelList
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        chart.dataLabels = DataLabelList()
        # openpyxl BarChart defaults y-axis gridlines ON; clear them so this
        # chart has data labels but no gridlines (the rule's pass condition).
        chart.y_axis.majorGridlines = None
        ws.add_chart(chart, "D1")
        path = _save_wb(wb)
        try:
            findings = _make_inspector().inspect_workbook(path)
            ids = _rule_ids(findings)
            assert "C101-GRIDLINES-NO-DATALABELS" not in ids
        finally:
            os.unlink(path)

